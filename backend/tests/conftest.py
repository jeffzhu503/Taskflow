# backend/tests/conftest.py
import io
import pytest
import docx
from fpdf import FPDF
import rag.task_generator as _task_generator_mod


@pytest.fixture(autouse=True)
def reset_task_generator_client():
    """Reset the singleton Anthropic client before each test for isolation."""
    _task_generator_mod._client = None
    yield
    _task_generator_mod._client = None


@pytest.fixture()
def docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Review Q2 budget variance with the finance team.")
    doc.add_paragraph("Schedule follow-up meeting with marketing.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def pdf_bytes() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(200, 10, "Review Q2 budget variance with the finance team.")
    pdf.ln()
    pdf.cell(200, 10, "Update headcount projections for H2.")
    return bytes(pdf.output())


@pytest.fixture()
def txt_bytes() -> bytes:
    return b"Review Q2 budget variance.\nUpdate headcount projections."
